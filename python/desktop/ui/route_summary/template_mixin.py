#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""RouteSummary RouteSummaryTemplateMixin."""
from __future__ import annotations

from typing import Optional, Dict, Any, List, Tuple
import html
import os
import sys
import re
import logging
import webbrowser
from functools import partial
from datetime import datetime, time as dt_time
from pathlib import Path

import pandas as pd
import openpyxl

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel,
    QTableWidget, QTableWidgetItem, QHeaderView, QGroupBox,
    QMessageBox, QFileDialog, QDateTimeEdit, QLineEdit,
    QTextEdit, QDoubleSpinBox, QSpinBox, QCheckBox, QComboBox,
    QDialog, QFormLayout, QDialogButtonBox, QTabWidget, QStyledItemDelegate, QStyle, QInputDialog,
    QSplitter, QApplication,
)
from PySide6.QtCore import Qt, QDateTime, QTime, Signal, QSettings, QUrl
from PySide6.QtGui import QColor, QShortcut, QKeySequence, QDrag, QGuiApplication, QBrush

from ui.star_rating_widget import StarRatingWidget

from .support import (
    WEBENGINE_AVAILABLE,
    QWebEngineView,
    ROUTE_SEGMENT_ROW_COLORS,
    COL_VISIT_INCLUDE,
    COL_VISIT_ORDER,
    COL_STORE_CODE,
    COL_STORE_NAME,
    COL_IN_TIME,
    COL_OUT_TIME,
    COL_STAY,
    COL_TRAVEL,
    COL_PROFIT,
    COL_QTY,
    COL_STAR,
    COL_NOTES,
    VISIT_TABLE_VISIBLE_COLUMNS,
    _VISIT_TABLE_DEFAULT_WIDTHS,
    _ROUTE_WORKFLOW_PIPELINE_SEGMENTS,
    _ROUTE_WORKFLOW_PIPELINE_SEP,
    _format_route_workflow_prefix_html,
    _format_route_workflow_pipeline_html,
    _template_include_from_db_value,
    _visit_include_checked,
    SafeInternalMoveTable,
    StoreSelectDialog,
    SavedRoutesDialog,
    RouteDatabase,
    StoreDatabase,
    RouteVisitDatabase,
    attach_table_column_width_persistence,
    reapply_table_column_widths,
    RouteMatchingService,
    CalculationService,
    TemplateGenerator,
    generate_route_map_urls,
    resolve_maps_api_key,
)



class RouteSummaryTemplateMixin:
    def load_template(
        self,
        file_path: Optional[str] = None,
        initial_dir: Optional[str] = None,
    ) -> Optional[str]:
        """
        テンプレートファイルを読み込む
        
        Args:
            file_path: 直接読み込むテンプレートファイルパス。
                       None の場合はファイル選択ダイアログを表示する。
            initial_dir: ファイル選択ダイアログの初期フォルダ（仕入CSV取込フォルダ等）。
        
        Returns:
            読み込んだファイルパス（キャンセル時や失敗時はNone）
        """
        try:
            # パスが指定されていない場合はダイアログで選択
            if not file_path:
                default_dir = ""
                if initial_dir and os.path.isdir(initial_dir):
                    default_dir = initial_dir
                else:
                    # 前回選択したテンプレートのフォルダ
                    last_path = self.settings.value("route_template/last_selected", "")
                    if last_path and os.path.exists(last_path):
                        default_dir = os.path.dirname(last_path)
                
                # ファイル選択ダイアログ
                file_path, _ = QFileDialog.getOpenFileName(
                    self,
                    "ルートテンプレートファイルを選択",
                    default_dir,
                    "Excelファイル (*.xlsx *.xls);;CSVファイル (*.csv);;すべてのファイル (*)"
                )
            
            # キャンセルされた場合
            if not file_path:
                return None
            
            # ファイルが存在するか確認
            if not os.path.exists(file_path):
                QMessageBox.warning(self, "警告", f"指定されたファイルが見つかりません:\n{file_path}")
                return None
            
            # ファイル拡張子で処理を分岐
            file_ext = os.path.splitext(file_path)[1].lower()
            
            if file_ext in ['.xlsx', '.xls']:
                # Excelファイルの読み込み
                self._load_excel_template(file_path)
            elif file_ext == '.csv':
                # CSVファイルの読み込み
                self._load_csv_template(file_path)
            else:
                QMessageBox.warning(self, "警告", "サポートされていないファイル形式です。")
                return None
            
            # 読み込み成功時はパスを保存
            self.last_loaded_template_path = file_path
            self.settings.setValue("route_template/last_selected", file_path)
            
            # 別ルートのテンプレを誤読みしたあと正しいテンプレで照合し直す場合、
            # 古い current_route_id のまま save すると DB の別行を更新してしまうため解除する
            self.current_route_id = None
            
            # 計算結果を更新
            getattr(self, 'update_calculation_results', lambda: None)()

            self._warn_route_template_time_issues_if_any()
            
            return file_path
            
        except Exception as e:
            QMessageBox.critical(self, "エラー", f"テンプレートの読み込みに失敗しました:\n{str(e)}")
            import traceback
            traceback.print_exc()
            return None

    def _collect_route_template_time_issues(self) -> List[str]:
        """読込済みの route_data・店舗訪問表から時刻の不整合を検出する。"""
        try:
            from desktop.services.route_template_time_validation import collect_route_template_time_issues
        except ImportError:
            from services.route_template_time_validation import collect_route_template_time_issues  # type: ignore

        route_data = getattr(self, "route_data", {}) or {}
        exclude = {"出発時刻", "帰宅時刻", "往路高速代", "復路高速代"}
        visits: List[Dict[str, Any]] = []
        for i in range(self.store_visits_table.rowCount()):
            code = self._get_table_item(i, COL_STORE_CODE)
            if code in exclude:
                continue
            visits.append(
                {
                    "store_code": code,
                    "store_name": self._get_table_item(i, COL_STORE_NAME),
                    "store_in_time": self._get_table_item(i, COL_IN_TIME),
                    "store_out_time": self._get_table_item(i, COL_OUT_TIME),
                }
            )
        return collect_route_template_time_issues(route_data, visits)

    def _warn_route_template_time_issues_if_any(self) -> None:
        """出発/帰宅・店舗IN/OUT が揃っていない場合に警告を表示する。"""
        issues = self._collect_route_template_time_issues()
        if not issues:
            return
        QMessageBox.warning(
            self,
            "ルートテンプレート: 時刻の確認",
            "次の時刻が対になっていないため、正しいデータが取得・計算できない可能性があります。\n\n"
            + "\n".join(issues)
            + "\n\n"
            "Excelの「店舗訪問詳細」シート（出発時刻・帰宅時刻の行と、各店舗の IN/OUT 列）を"
            "修正してから、再度「ルートテンプレ読込」を実行してください。",
        )

    def _parse_route_name_from_template_filename(self, file_path: str) -> str:
        """route_template_ルート名_YYYYMMDD.xlsx 形式からルート名を抽出。"""
        stem = Path(file_path).stem
        m = re.match(r"route_template_(.+)_(\d{8})$", stem, re.IGNORECASE)
        if m:
            return m.group(1).strip()
        return ""

    def _set_route_name_from_template(self, route_name: str, file_path: str = "") -> None:
        """テンプレートから読んだルート名を route_data とコンボボックスに反映。"""
        route_name = (route_name or "").strip()
        if not route_name and file_path:
            route_name = self._parse_route_name_from_template_filename(file_path)
        if not route_name:
            return
        if not hasattr(self, "route_data") or not self.route_data:
            self.route_data = {}
        self.route_data["route_name"] = route_name
        self.route_data["route_name_display"] = route_name
        self.update_route_codes()
        idx = self.route_code_combo.findText(route_name)
        if idx >= 0:
            self.route_code_combo.setCurrentIndex(idx)
        else:
            self.route_code_combo.addItem(route_name)
            self.route_code_combo.setCurrentText(route_name)

    def _load_excel_template(self, file_path: str):
        """Excelテンプレートファイルを読み込む"""
        try:
            import openpyxl
            
            wb = openpyxl.load_workbook(file_path, data_only=True)
            
            # 店舗訪問詳細シートを探す
            visit_sheet = None
            for sheet_name in wb.sheetnames:
                if '店舗訪問詳細' in sheet_name or '訪問' in sheet_name:
                    visit_sheet = wb[sheet_name]
                    break
            
            if not visit_sheet:
                # デフォルトで最初のシートを使用
                visit_sheet = wb.active
            
            # ルート日付とルートコードを取得（1-2行目）
            route_date = None
            route_code = None
            
            # 1行目: 日付
            date_label = visit_sheet.cell(row=1, column=1).value
            if date_label and ('日付' in str(date_label) or 'date' in str(date_label).lower()):
                date_value = visit_sheet.cell(row=1, column=2).value
                if date_value:
                    if isinstance(date_value, datetime):
                        route_date = date_value.strftime('%Y-%m-%d')
                    else:
                        route_date = str(date_value)
            
            # 2行目: ルート（表示名＝所属ルート名）
            route_label = visit_sheet.cell(row=2, column=1).value
            route_name_from_file = ""
            if route_label and ('ルート' in str(route_label) or 'route' in str(route_label).lower()):
                route_value = visit_sheet.cell(row=2, column=2).value
                if route_value:
                    route_name_from_file = str(route_value).strip()
            
            # ルート情報を設定
            if not hasattr(self, 'route_data') or not self.route_data:
                self.route_data = {}
            if route_date:
                self.route_data['route_date'] = route_date
                try:
                    if isinstance(route_date, str):
                        route_date_obj = datetime.strptime(route_date, '%Y-%m-%d').date()
                    else:
                        route_date_obj = route_date
                    self.route_date_edit.setDate(route_date_obj)
                except Exception:
                    pass
            
            self._set_route_name_from_template(route_name_from_file, file_path)
            
            # 店舗訪問詳細を読み込み（3行目がヘッダー、4行目以降がデータ）
            visits = []
            header_row = 3
            data_start_row = 4
            
            # ヘッダーを取得
            headers = []
            for col in range(1, visit_sheet.max_column + 1):
                header_value = visit_sheet.cell(row=header_row, column=col).value
                if header_value:
                    headers.append(str(header_value).strip())
                else:
                    headers.append('')
            
            # 出発時刻・帰宅時刻・往路高速代・復路高速代を保存する変数
            departure_time = None
            return_time = None
            toll_fee_outbound = 0
            toll_fee_return = 0
            
            # データ行を読み込み
            for row in range(data_start_row, visit_sheet.max_row + 1):
                # 店舗コードが空の行はスキップ
                store_code_cell = visit_sheet.cell(row=row, column=1)
                store_code = store_code_cell.value
                if not store_code or str(store_code).strip() == '':
                    continue
                
                # 除外対象の店舗コードをチェック（値を取得してからスキップ）
                store_code_str = str(store_code).strip()
                if store_code_str == '出発時刻':
                    # 出発時刻を取得（B列）
                    departure_time = self._format_time_value(visit_sheet.cell(row=row, column=2).value)
                    continue
                elif store_code_str == '帰宅時刻':
                    # 帰宅時刻を取得（B列）
                    return_time = self._format_time_value(visit_sheet.cell(row=row, column=2).value)
                    continue
                elif store_code_str == '往路高速代':
                    # 往路高速代を取得（B列）
                    toll_value = visit_sheet.cell(row=row, column=2).value
                    toll_fee_outbound = self._safe_float(str(toll_value or '')) or 0
                    continue
                elif store_code_str == '復路高速代':
                    # 復路高速代を取得（B列）
                    toll_value = visit_sheet.cell(row=row, column=2).value
                    toll_fee_return = self._safe_float(str(toll_value or '')) or 0
                    continue
                
                visit = {
                    'visit_order': len(visits) + 1,
                    'store_code': store_code_str,
                    'store_name': str(visit_sheet.cell(row=row, column=2).value or '').strip(),
                    'store_in_time': self._format_time_value(visit_sheet.cell(row=row, column=3).value),
                    'store_out_time': self._format_time_value(visit_sheet.cell(row=row, column=4).value),
                    'stay_duration': self._safe_float(str(visit_sheet.cell(row=row, column=5).value or '')) or 0.0,
                    'store_notes': str(visit_sheet.cell(row=row, column=6).value or '').strip(),
                }
                visits.append(visit)
            
            # 読み込んだ値をroute_dataに保存
            if departure_time or return_time or toll_fee_outbound or toll_fee_return:
                if not hasattr(self, 'route_data') or not self.route_data:
                    self.route_data = {}
                
                if departure_time:
                    # ルート日付と結合してdatetime形式に
                    if route_date:
                        self.route_data['departure_time'] = f"{route_date} {departure_time}:00"
                    else:
                        self.route_data['departure_time'] = departure_time
                
                if return_time:
                    # ルート日付と結合してdatetime形式に
                    if route_date:
                        self.route_data['return_time'] = f"{route_date} {return_time}:00"
                    else:
                        self.route_data['return_time'] = return_time
                
                if toll_fee_outbound:
                    self.route_data['toll_fee_outbound'] = toll_fee_outbound
                
                if toll_fee_return:
                    self.route_data['toll_fee_return'] = toll_fee_return
            
            # テーブルに反映
            self.store_visits_table.blockSignals(True)
            try:
                self.store_visits_table.setRowCount(len(visits))
                for i, visit in enumerate(visits):
                    self._fill_visit_table_row(
                        i,
                        store_code=visit.get('store_code', ''),
                        store_name=visit.get('store_name', ''),
                        in_time=visit.get('store_in_time', ''),
                        out_time=visit.get('store_out_time', ''),
                        notes=visit.get('store_notes', ''),
                        include_checked=True,
                    )
            finally:
                self.store_visits_table.blockSignals(False)
            
            # 滞在時間と移動時間を自動計算
            self.recalc_travel_times()
            
            # 訪問順序を更新
            self.update_visit_order()
            
            QMessageBox.information(self, "完了", f"テンプレートを読み込みました。\n店舗数: {len(visits)}件")
            
        except Exception as e:
            raise Exception(f"Excelファイルの読み込みエラー: {str(e)}")

    def _load_csv_template(self, file_path: str):
        """CSVテンプレートファイルを読み込む"""
        try:
            # CSVファイルの読み込み（UTF-8 BOM対応）
            df = pd.read_csv(file_path, encoding='utf-8-sig', dtype=str, keep_default_na=False)
            
            # ルート情報セクションと店舗訪問詳細セクションを分離
            route_data = {}
            visits = []
            
            in_route_section = False
            in_visit_section = False
            
            with open(file_path, 'r', encoding='utf-8-sig') as f:
                lines = f.readlines()
            
            for line in lines:
                line = line.strip()
                if not line or line.startswith('#'):
                    continue
                
                if 'ルート情報' in line:
                    in_route_section = True
                    in_visit_section = False
                    continue
                elif '店舗訪問詳細' in line:
                    in_route_section = False
                    in_visit_section = True
                    continue
                
                if in_route_section:
                    # ルート情報の処理（簡易実装）
                    parts = line.split(',')
                    if len(parts) >= 2:
                        key = parts[0].strip()
                        value = parts[1].strip() if len(parts) > 1 else ''
                        if key == 'ルート日付':
                            route_data['route_date'] = value
                        elif key in ('ルートコード', 'ルート名', 'ルート'):
                            route_data['route_name'] = value
                            route_data['route_code'] = value
                        elif key == '出発時間':
                            route_data['departure_time'] = value
                        elif key == '帰宅時間':
                            route_data['return_time'] = value
                        elif key == '往路高速代':
                            try:
                                route_data['toll_fee_outbound'] = float(value) if value else 0
                            except (ValueError, TypeError):
                                route_data['toll_fee_outbound'] = 0
                        elif key == '復路高速代':
                            try:
                                route_data['toll_fee_return'] = float(value) if value else 0
                            except (ValueError, TypeError):
                                route_data['toll_fee_return'] = 0
                elif in_visit_section:
                    # 店舗訪問詳細の処理
                    parts = line.split(',')
                    if len(parts) >= 2 and parts[0].strip() != '訪問順序':
                        visit = {
                            'visit_order': len(visits) + 1,
                            'store_code': parts[1].strip() if len(parts) > 1 else '',
                            'store_name': parts[2].strip() if len(parts) > 2 else '',
                            'store_in_time': parts[3].strip() if len(parts) > 3 else '',
                            'store_out_time': parts[4].strip() if len(parts) > 4 else '',
                            'store_notes': parts[11].strip() if len(parts) > 11 else '',
                        }
                        if visit['store_code']:
                            visits.append(visit)
            
            # ルート情報を設定
            route_date_str = None
            if 'route_date' in route_data:
                route_date_str = route_data['route_date']
                try:
                    route_date_obj = datetime.strptime(route_date_str, '%Y-%m-%d').date()
                    self.route_date_edit.setDate(route_date_obj)
                except Exception:
                    pass
            
            if not hasattr(self, 'route_data') or not self.route_data:
                self.route_data = {}

            route_name_from_file = (
                str(route_data.get('route_name') or route_data.get('route_code') or '').strip()
            )
            self._set_route_name_from_template(route_name_from_file, file_path)
            if route_date_str:
                self.route_data['route_date'] = route_date_str
            
            if 'departure_time' in route_data:
                dep_time = route_data['departure_time']
                if route_date_str:
                    self.route_data['departure_time'] = f"{route_date_str} {dep_time}:00" if ' ' not in str(dep_time) else dep_time
                else:
                    self.route_data['departure_time'] = dep_time
            
            if 'return_time' in route_data:
                ret_time = route_data['return_time']
                if route_date_str:
                    self.route_data['return_time'] = f"{route_date_str} {ret_time}:00" if ' ' not in str(ret_time) else ret_time
                else:
                    self.route_data['return_time'] = ret_time
            
            if 'toll_fee_outbound' in route_data:
                self.route_data['toll_fee_outbound'] = route_data['toll_fee_outbound']
            
            if 'toll_fee_return' in route_data:
                self.route_data['toll_fee_return'] = route_data['toll_fee_return']
            
            # テーブルに反映（Excelと同じ処理）
            self.store_visits_table.blockSignals(True)
            try:
                self.store_visits_table.setRowCount(len(visits))
                for i, visit in enumerate(visits):
                    self._fill_visit_table_row(
                        i,
                        store_code=visit.get('store_code', ''),
                        store_name=visit.get('store_name', ''),
                        in_time=visit.get('store_in_time', ''),
                        out_time=visit.get('store_out_time', ''),
                        notes=visit.get('store_notes', ''),
                        include_checked=True,
                    )
            finally:
                self.store_visits_table.blockSignals(False)
            
            self.recalc_travel_times()
            self.update_visit_order()
            
            QMessageBox.information(self, "完了", f"テンプレートを読み込みました。\n店舗数: {len(visits)}件")
            
        except Exception as e:
            raise Exception(f"CSVファイルの読み込みエラー: {str(e)}")

    def _format_time_value(self, value) -> str:
        """時刻値を文字列に変換"""
        if value is None:
            return ''
        if isinstance(value, datetime):
            return value.strftime('%H:%M')
        if isinstance(value, dt_time):
            return value.strftime('%H:%M')
        value_str = str(value).strip()
        if ':' in value_str:
            # HH:MM形式の場合はそのまま返す
            parts = value_str.split(':')
            if len(parts) >= 2:
                return f"{parts[0].zfill(2)}:{parts[1].zfill(2)}"
        return value_str

    def generate_template(self):
        """テンプレート生成

        デフォルトフォルダ直下に
            YYYYMMDDルート名
        というフォルダを自動作成し、
        その中にテンプレートファイルと
        「商品画像」「レシート画像」フォルダを作成する。
        """
        try:
            from pathlib import Path

            route_code = self.get_selected_route_code()
            if not route_code:
                QMessageBox.warning(self, "エラー", "ルートコードを選択してください。")
                return None

            # ルート名（日本語名）を取得（なければコードをそのまま使用）
            route_name = self.store_db.get_route_name_by_code(route_code) or route_code

            # ルート日付（YYYYMMDD）を取得
            route_qdate = self.route_date_edit.date()
            route_date_str = route_qdate.toString("yyyyMMdd")

            # デフォルトフォルダが未設定の場合は、ユーザーに選択してもらう
            base_dir = self.template_save_default_dir if self.template_save_default_dir and os.path.isdir(self.template_save_default_dir) else ""
            if not base_dir:
                base_dir = QFileDialog.getExistingDirectory(
                    self,
                    "テンプレート保存用のデフォルトフォルダを選択",
                    str(Path.home()),
                )
                if not base_dir:
                    return None
                self.update_template_save_default_dir(base_dir)

            base_path = Path(base_dir)

            # フォルダ名: YYYYMMDDルート名（ファイル名に使えない文字は簡易的に置換）
            unsafe_chars = '\\/:*?"<>|'
            safe_route_name = "".join("_" if ch in unsafe_chars else ch for ch in route_name.strip())
            folder_name = f"{route_date_str}{safe_route_name}"
            route_folder = base_path / folder_name

            # ルートフォルダとサブフォルダを作成
            route_folder.mkdir(parents=True, exist_ok=True)
            (route_folder / "商品画像").mkdir(exist_ok=True)
            (route_folder / "レシート画像").mkdir(exist_ok=True)

            # テンプレートファイルパス（Excel固定）
            default_filename = f"route_template_{safe_route_name}_{route_date_str}.xlsx"
            file_path = str(route_folder / default_filename)
            
            # 選択されたルートの店舗一覧を取得
            stores = []
            store_codes = []
            if route_name:
                # テーブル備考を店舗マスタへ上書き（次回読み込み・店舗一覧タブと同期）
                if self.store_visits_table.rowCount() > 0:
                    notes_updated, notes_skipped = self._persist_visit_table_notes_to_store_master()
                    print(
                        f"店舗マスタ備考を上書き: {notes_updated}件"
                        + (f"（スキップ {notes_skipped}件）" if notes_skipped else "")
                    )
                # テーブルから訪問順序を取得（テーブルにデータがある場合）
                table_stores = self.get_stores_from_table(for_template_output=True)
                if table_stores:
                    # テーブルの訪問順序を使用（チェック済み店舗のみ）
                    stores = table_stores
                    store_codes = [
                        (store.get('store_code') or store.get('supplier_code'))
                        for store in stores
                        if store.get('store_code') or store.get('supplier_code')
                    ]
                elif self.store_visits_table.rowCount() > 0:
                    QMessageBox.warning(
                        self,
                        "警告",
                        "テンプレートに出力する店舗が選択されていません。\n"
                        "「出力」列のチェックを確認してください。",
                    )
                    return
                else:
                    # テーブルにデータがない場合はデータベースから取得（template_include=1 のみ）
                    stores = [
                        store for store in self.get_stores_for_route(route_name)
                        if _template_include_from_db_value(store.get('template_include'))
                    ]
                    for store in stores:
                        any_code = self._visit_store_code_from_dict(store)
                        if any_code:
                            store_info = self.store_db.get_store_by_code(any_code)
                            if store_info:
                                store['notes'] = self._resolve_store_master_notes(store_info)
                    store_codes = [
                        (store.get('store_code') or store.get('supplier_code'))
                        for store in stores
                        if store.get('store_code') or store.get('supplier_code')
                    ]
            
            if not store_codes:
                QMessageBox.warning(self, "警告", "テンプレートに出力する店舗がありません。")
                return
            
            if not TemplateGenerator:
                QMessageBox.warning(self, "エラー", "テンプレート生成機能が利用できません")
                return
            
            # ルート日付を取得（QDateをdatetime.dateに変換）
            try:
                qdate = self.route_date_edit.dateTime().date()
                # QDateをdatetime.dateに変換
                route_date = datetime(qdate.year(), qdate.month(), qdate.day()).date()
                print(f"ルート日付取得: {route_date} (QDate: {qdate.year()}-{qdate.month()}-{qdate.day()})")
            except Exception as e:
                print(f"ルート日付取得エラー: {e}")
                import traceback
                print(traceback.format_exc())
                route_date = None
            
            # Excelテンプレート生成（ルート名とルート日付を渡す）
            print(f"Excelテンプレート生成開始: ファイル={file_path}, ルート名={route_name}, 店舗数={len(stores) if stores else len(store_codes) if store_codes else 0}, ルート日付={route_date}")
            success = TemplateGenerator.generate_excel_template(file_path, route_name, store_codes, stores, route_date)
            
            if success:
                QMessageBox.information(
                    self,
                    "成功",
                    f"テンプレートを生成しました:\n{file_path}\n\n"
                    f"以下のフォルダも自動作成しました:\n"
                    f"  - {route_folder}\n"
                    f"  - {route_folder / '商品画像'}\n"
                    f"  - {route_folder / 'レシート画像'}",
                )
            else:
                QMessageBox.warning(self, "エラー", "テンプレートの生成に失敗しました。\n詳細はコンソールを確認してください。")
                
        except Exception as e:
            import traceback
            error_detail = traceback.format_exc()
            print(f"テンプレート生成エラー詳細:\n{error_detail}")
            QMessageBox.critical(self, "エラー", f"テンプレート生成中にエラーが発生しました:\n{str(e)}\n\n詳細はコンソールを確認してください。")

    def set_template_root_directory(self):
        """テンプレート保存用のデフォルト（起点）フォルダを設定"""
        current_dir = self.template_save_default_dir if self.template_save_default_dir and os.path.isdir(self.template_save_default_dir) else str(Path.home())
        selected_dir = QFileDialog.getExistingDirectory(
            self,
            "テンプレート保存用のデフォルトフォルダを選択",
            current_dir,
        )
        if not selected_dir:
            return

        self.update_template_save_default_dir(selected_dir)
        QMessageBox.information(
            self,
            "デフォルトフォルダ設定",
            f"ルートテンプレートを保存する起点フォルダを次の場所に設定しました:\n{selected_dir}",
        )

    def on_template_save_dir_edit_finished(self):
        """手入力でテンプレート保存フォルダを更新"""
        text = self.template_save_dir_edit.text().strip()
        if not text:
            self.update_template_save_default_dir("")
            return
        expanded = os.path.expanduser(text)
        if not os.path.isdir(expanded):
            QMessageBox.warning(self, "エラー", "指定したフォルダが存在しません。")
            # 元の値へ戻す
            self.template_save_dir_edit.blockSignals(True)
            self.template_save_dir_edit.setText(self.template_save_default_dir)
            self.template_save_dir_edit.blockSignals(False)
            return
        self.update_template_save_default_dir(os.path.abspath(expanded))

    def update_template_save_default_dir(self, directory: str):
        """テンプレート保存フォルダの設定を保存"""
        normalized = directory if directory else ""
        if normalized and not os.path.isdir(normalized):
            return
        self.template_save_default_dir = normalized
        if hasattr(self, "template_save_dir_edit"):
            self.template_save_dir_edit.blockSignals(True)
            self.template_save_dir_edit.setText(self.template_save_default_dir)
            self.template_save_dir_edit.blockSignals(False)
        if self.settings is not None:
            self.settings.setValue("route_template/default_save_dir", self.template_save_default_dir)

