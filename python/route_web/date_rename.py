# -*- coding: utf-8 -*-
"""ルート日付の変更と、箱・画像・Excel のファイル名の付け替え。"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional

from route_web.registry import update_folder_location

_ISO_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_YMD_RE = re.compile(r"^\d{8}$")


class RouteDateRenameError(Exception):
    """日付の変更先に、同じ名前のフォルダやファイルがすでにある。"""


def parse_route_date(value: str) -> str:
    """日付を YYYY-MM-DD にする。不正なら ValueError。"""
    text = (value or "").strip().replace("/", "-").replace(".", "-")
    if _YMD_RE.fullmatch(text):
        text = f"{text[:4]}-{text[4:6]}-{text[6:8]}"
    if not _ISO_RE.fullmatch(text):
        raise ValueError("日付は 2026-10-17 の形式で選んでください。")
    try:
        datetime.strptime(text, "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError("存在しない日付です。") from exc
    return text


def _ymd(iso: str) -> str:
    return iso.replace("-", "")


def replace_route_date_text(text: str, old_iso: str, new_iso: str) -> str:
    """ファイル名や文章の中のルート日付だけを付け替える。"""
    if not text or old_iso == new_iso:
        return text
    updated = text.replace(old_iso, new_iso)
    old_ymd = _ymd(old_iso)
    new_ymd = _ymd(new_iso)
    return re.sub(rf"(?<!\d){old_ymd}(?!\d)", new_ymd, updated)


def renamed_directory_name(name: str, old_iso: str, new_iso: str) -> str:
    """YYYYMMDD名前 と 名前YYYYMMDD の両方を新しい日付にする。"""
    old_ymd = _ymd(old_iso)
    new_ymd = _ymd(new_iso)
    if name.startswith(old_ymd):
        return new_ymd + name[len(old_ymd) :]
    if name.endswith(old_ymd):
        return name[: -len(old_ymd)] + new_ymd
    return replace_route_date_text(name, old_iso, new_iso)


def _rewrite_tree(value: Any, old_iso: str, new_iso: str) -> Any:
    if isinstance(value, str):
        return replace_route_date_text(value, old_iso, new_iso)
    if isinstance(value, list):
        return [_rewrite_tree(item, old_iso, new_iso) for item in value]
    if isinstance(value, dict):
        return {key: _rewrite_tree(item, old_iso, new_iso) for key, item in value.items()}
    return value


def _planned_file_renames(root: Path, old_iso: str, new_iso: str) -> list[tuple[Path, Path]]:
    planned: list[tuple[Path, Path]] = []
    if not root.is_dir():
        return planned
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        new_name = replace_route_date_text(path.name, old_iso, new_iso)
        if new_name == path.name:
            continue
        dest = path.with_name(new_name)
        planned.append((path, dest))
    dests = [dest for _src, dest in planned]
    if len(dests) != len(set(dests)):
        raise RouteDateRenameError("日付を変えると、同じ名前になるファイルが箱の中にあります。")
    for _src, dest in planned:
        if dest.exists():
            raise RouteDateRenameError(f"同じ名前のファイルがすでにあります: {dest.name}")
    return planned


def _rename_files(root: Path, old_iso: str, new_iso: str) -> None:
    planned = _planned_file_renames(root, old_iso, new_iso)
    # 深いパスから変えて、親を先に触らない。
    planned.sort(key=lambda pair: len(pair[0].parts), reverse=True)
    for src, dest in planned:
        src.rename(dest)


def _rewrite_sidecar_json(root: Path, old_iso: str, new_iso: str) -> None:
    if not root.is_dir():
        return
    for path in root.rglob("*.json"):
        if path.name == "route.json":
            continue
        try:
            raw = path.read_text(encoding="utf-8")
        except (OSError, UnicodeError):
            continue
        updated = replace_route_date_text(raw, old_iso, new_iso)
        if updated != raw:
            path.write_text(updated, encoding="utf-8")


def _update_excel_date(path: Path, new_iso: str) -> None:
    """Excel 1行目の日付セルを新しい日付にする。失敗してもファイル名の変更は残す。"""
    if path.suffix.lower() not in {".xlsx", ".xlsm"} or not path.is_file():
        return
    try:
        import openpyxl
    except ImportError:
        return
    try:
        book = openpyxl.load_workbook(path)
    except Exception:
        return
    try:
        sheet = book.worksheets[0]
        label = sheet.cell(row=1, column=1).value
        if not label or "日付" not in str(label):
            return
        sheet.cell(row=1, column=2).value = datetime.strptime(new_iso, "%Y-%m-%d").date()
        sheet.cell(row=1, column=2).number_format = "yyyy/mm/dd"
        book.save(path)
    except Exception:
        return
    finally:
        book.close()


def _rename_directory(src: Path, new_name: str) -> Path:
    if new_name == src.name:
        return src
    dest = src.parent / new_name
    if dest.exists():
        raise RouteDateRenameError(f"同じ名前のフォルダがすでにあります: {dest.name}")
    src.rename(dest)
    return dest


def _insurance_dir(folder: Path, insurance_root: Optional[Path]) -> Optional[Path]:
    if insurance_root is not None:
        candidate = Path(insurance_root) / folder.name
        return candidate if candidate.is_dir() else None
    try:
        from services.route_folder_import import matching_insurance_dir
    except ImportError:
        from desktop.services.route_folder_import import matching_insurance_dir  # type: ignore
    return matching_insurance_dir(folder)


def change_route_date(
    folder: Path,
    doc: Dict[str, Any],
    old_iso: str,
    new_iso: str,
    *,
    insurance_root: Optional[Path] = None,
) -> Dict[str, Any]:
    """箱の中の日付入りファイルと、箱・ルート保険のフォルダ名を付け替える。"""
    root = Path(folder)
    if not root.is_dir():
        raise RouteDateRenameError("ルートのフォルダが見つかりません。")
    old_iso = parse_route_date(old_iso)
    new_iso = parse_route_date(new_iso)
    if old_iso == new_iso:
        return doc

    new_dir_name = renamed_directory_name(root.name, old_iso, new_iso)
    if new_dir_name != root.name and (root.parent / new_dir_name).exists():
        raise RouteDateRenameError(f"同じ名前のフォルダがすでにあります: {new_dir_name}")

    insurance = _insurance_dir(root, insurance_root)
    if insurance is not None:
        insurance_name = renamed_directory_name(insurance.name, old_iso, new_iso)
        if insurance_name != insurance.name and (insurance.parent / insurance_name).exists():
            raise RouteDateRenameError(f"ルート保険に同じ名前のフォルダがあります: {insurance_name}")
        _planned_file_renames(insurance, old_iso, new_iso)

    _planned_file_renames(root, old_iso, new_iso)
    _rename_files(root, old_iso, new_iso)
    _rewrite_sidecar_json(root, old_iso, new_iso)
    for excel in list(root.rglob("*.xlsx")) + list(root.rglob("*.xlsm")):
        _update_excel_date(excel, new_iso)

    if insurance is not None:
        _rename_files(insurance, old_iso, new_iso)
        _rewrite_sidecar_json(insurance, old_iso, new_iso)
        for excel in list(insurance.rglob("*.xlsx")) + list(insurance.rglob("*.xlsm")):
            _update_excel_date(excel, new_iso)
        insurance = _rename_directory(insurance, renamed_directory_name(insurance.name, old_iso, new_iso))

    old_folder = root
    new_folder = _rename_directory(root, new_dir_name)
    updated = _rewrite_tree(doc, old_iso, new_iso)
    if not isinstance(updated, dict):
        updated = dict(doc)
    updated["route_date"] = new_iso
    updated["folder_path"] = str(new_folder)
    update_folder_location(old_folder, new_folder, new_iso)
    return updated
